from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_ROOT = PROJECT_ROOT / "tests" / "testcases"

SUITE_FILES = {
    "smoke": [
        PROJECT_ROOT / "tests" / "smoke" / "LOGIN_SMOKE_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "smoke" / "ADD_MEMBER_SMOKE_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "smoke" / "MANAGE_MEMBERS_SMOKE_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "smoke" / "CREATE_TEAM_SMOKE_TESTSUITE.md",
    ],
    "sanity": [
        PROJECT_ROOT / "tests" / "sanity" / "LOGIN_PAGE_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "sanity" / "ADD_MEMBER_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "sanity" / "MANAGE_MEMBERS_TESTSUITE.md",
        PROJECT_ROOT / "tests" / "sanity" / "CREATE_TEAM_TESTSUITE.md",
    ],
}

HEADER_FILL = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADERS = ["Module", "Suite", "Case ID", "Test Case", "Steps", "Expected Result", "Status", "Remarks"]


def normalize_module_name(title: str) -> str:
    cleaned = title.replace("Test Suite", "").replace("Test Case Suite", "").strip()
    cleaned = cleaned.replace("ProtonNxt ", "").replace(" - ", " ")
    lowered = cleaned.lower()

    if "login" in lowered:
        return "Login"
    if "add member" in lowered or "add team member" in lowered:
        return "Add Member"
    if "manage members" in lowered:
        return "Manage Members"
    if "create team" in lowered:
        return "Create Team"

    return cleaned


def parse_suite_markdown(path: Path):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].lstrip("#").strip() if lines else path.stem
    module_name = normalize_module_name(title)

    cases = []
    current = None
    section = None

    for raw_line in lines[1:]:
        line = raw_line.strip()

        match = re.match(r"^##\s+(.+?)\s+-\s+(.+)$", line)
        if match:
            if current:
                cases.append(current)
            current = {
                "case_id": match.group(1).strip(),
                "test_case": match.group(2).strip(),
                "steps": [],
                "expected": [],
            }
            section = None
            continue

        if current is None:
            continue

        if line == "**Steps**":
            section = "steps"
            continue

        if line == "**Expected Result**":
            section = "expected"
            continue

        if line == "---":
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)

        if section == "steps" and numbered:
            current["steps"].append(numbered.group(1).strip())
        elif section == "expected" and bullet:
            current["expected"].append(bullet.group(1).strip())

    if current:
        cases.append(current)

    return module_name, cases


def autosize_columns(worksheet):
    for column_cells in worksheet.columns:
        values = [str(cell.value or "") for cell in column_cells]
        max_length = max(len(value) for value in values)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(max_length + 2, 14),
            60,
        )


def write_workbook(module_name: str, suite_name: str, cases: list[dict], output_path: Path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test Cases"
    worksheet.append(HEADERS)

    for case in cases:
        worksheet.append(
            [
                module_name,
                suite_name.title(),
                case["case_id"],
                case["test_case"],
                " ".join(f"{index + 1}. {step}" for index, step in enumerate(case["steps"])),
                "; ".join(case["expected"]),
                "Not Run",
                "",
            ]
        )

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_columns(worksheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main():
    module_cases: dict[str, list[dict]] = {}

    for suite_name, file_paths in SUITE_FILES.items():
        for path in file_paths:
            if not path.exists():
                print(f"Skipped missing suite file: {path}")
                continue
            module_name, cases = parse_suite_markdown(path)
            output_name = f"{path.stem.lower()}.xlsx"
            output_path = OUTPUT_ROOT / suite_name / output_name
            write_workbook(module_name, suite_name, cases, output_path)
            print(f"Created: {output_path}")

            module_cases.setdefault(module_name, [])
            for case in cases:
                module_cases[module_name].append(
                    {
                        **case,
                        "suite": suite_name.title(),
                    }
                )

    for module_name, cases in module_cases.items():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Test Cases"
        worksheet.append(HEADERS)

        for case in cases:
            worksheet.append(
                [
                    module_name,
                    case["suite"],
                    case["case_id"],
                    case["test_case"],
                    " ".join(f"{index + 1}. {step}" for index, step in enumerate(case["steps"])),
                    "; ".join(case["expected"]),
                    "Not Run",
                    "",
                ]
            )

        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        autosize_columns(worksheet)

        combined_name = module_name.lower().replace(" ", "_") + "_all_testcases.xlsx"
        combined_path = OUTPUT_ROOT / "module_wise" / combined_name
        combined_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(combined_path)
        print(f"Created: {combined_path}")


if __name__ == "__main__":
    main()
