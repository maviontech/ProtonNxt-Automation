from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
CSV_ROOT = PROJECT_ROOT / "testcase_files"
XLSX_ROOT = CSV_ROOT / "xlsx"


HEADER_FILL = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def autosize_columns(worksheet):
    for column_cells in worksheet.columns:
        values = [str(cell.value or "") for cell in column_cells]
        max_length = max(len(value) for value in values)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(max_length + 2, 14),
            60,
        )


def convert_csv_to_xlsx(csv_path: Path, xlsx_path: Path):
    with csv_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = csv_path.stem[:31]

    for row in rows:
        worksheet.append(row)

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_columns(worksheet)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def main():
    for csv_path in CSV_ROOT.rglob("*.csv"):
        relative = csv_path.relative_to(CSV_ROOT)
        xlsx_path = XLSX_ROOT / relative.with_suffix(".xlsx")
        convert_csv_to_xlsx(csv_path, xlsx_path)
        print(f"Created: {xlsx_path}")


if __name__ == "__main__":
    main()
